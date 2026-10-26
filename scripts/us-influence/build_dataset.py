from pathlib import Path
from bs4 import BeautifulSoup
from collections import Counter,defaultdict
import openpyxl,pycountry,json,re,datetime,hashlib,csv,zipfile,io,unicodedata

ROOT=Path(__file__).parent; RAW=ROOT/'raw'; OUT=ROOT/'parsed'
DEST=ROOT.parents[1]/'public'/'us-influence'/'data'; DEST.mkdir(parents=True,exist_ok=True)
for old in DEST.glob('*.json'):old.unlink()
def norm(s):return re.sub(r'[^a-z0-9]','',unicodedata.normalize('NFKD',str(s)).encode('ascii','ignore').decode().lower())
def write(p,x):p.write_text(json.dumps(x,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
def date(v):return v.date().isoformat() if isinstance(v,datetime.datetime) else None
def number(v):
 if isinstance(v,(int,float)):return v
 if isinstance(v,str) and re.fullmatch(r'-?[\d,]+(?:\.\d+)?',v.strip()):return float(v.replace(',',''))
 return None
manifests=[]
for f in ROOT.glob('manifest*.json'):manifests.extend(json.loads(f.read_text(encoding='utf-8')))
mf={m['file']:m for m in manifests if 'error' not in m};sources={}
def source(file,kind,title=None,coverage=None):
 sid=re.sub(r'[^a-zA-Z0-9_-]','-',file.rsplit('.',1)[0]);m=mf[file]
 sources[sid]={'id':sid,'title':title or file,'kind':kind,'url':m['url'],'retrieved':'2026-10-26','sha256':m['sha256'],'coverage':coverage,'file':file}
 return sid
names={};entities={}
for c in pycountry.countries:
 entities[c.alpha_2]={'id':c.alpha_2,'name':c.name,'kind':'territory','iso3':c.alpha_3,'aliases':[]}
 for n in [c.name,getattr(c,'official_name',''),getattr(c,'common_name',''),c.alpha_2,c.alpha_3]:
  if n:names[norm(n)]=c.alpha_2
aliases={
 'China (the People\'s Republic of)':'CN','Gambia (Republic of The)':'GM','Naoero':'NR','Netherlands (Kingdom of the)':'NL',
 'Antigua & Barbuda':'AG','Cayman Islands (The)':'KY','Ceylon':'LK','Congo (Kinshasa)':'CD','Congo (Kinshasa) (Zaire)':'CD','Congo (The)':'CG',
 'Kosova':'XK','Macau':'MO','Marshal Islands (The)':'MH','Myanmar (Burma)':'MM','Papua and New Guinea':'PG',
 'Sint Maarten':'SX','Somali Democratic Republic':'SO','St. Barthelemy':'BL','St. Barts':'BL','St. Christopher (St. Kitts) & Nevis':'KN',
 'St. Martin':'MF','St. Pierre and Miquelon':'PM','St. Vincent & The Grenadines':'VC','Sudan (The)':'SD','The Netherlands':'NL',
 'Timor-Leste (East Timor)':'TL','Trinidad & Tobago':'TT','Turkey (Turkiye)':'TR','Western Samoa':'WS',
 'Nagorno Karabagh':'extra-nagorno-karabakh','Nagorno Karabakh':'extra-nagorno-karabakh',
 'Great Britain':'GB','Britain':'GB',
 'QATAR':'QA','Korea, South':'KR','Korea, North':'KP','South Korea':'KR','North Korea':'KP','Korea, Republic of':'KR','Republic of Korea':'KR',
 'Democratic Republic of the Congo':'CD','Congo, Democratic Republic of the':'CD','Congo, Republic of the':'CG','Congo (Brazzaville)':'CG',
 'Bahamas, The':'BS','The Bahamas':'BS','Bahamas (The)':'BS','Gambia (The)':'GM','Congo (Democratic Republic of the)':'CD',
 'Congo, Democratic Republic of':'CD','Congo, Republic of':'CG',
 'Cote d\'Ivoire (Ivory Coast)':'CI','Ivory Coast':'CI','Cote D’Ivoire':'CI','Côte D’Ivoire':'CI','Czech Republic':'CZ','Czechia':'CZ',
 'Burma':'MM','Türkiye':'TR','Turkey':'TR','Turkiye':'TR','Vietnam':'VN','Viet Nam':'VN','Laos':'LA',
 'Russia':'RU','Russian Federation':'RU','Iran':'IR','Syria':'SY','Tanzania':'TZ','Venezuela':'VE','Bolivia':'BO',
 'Micronesia':'FM','Moldova':'MD','Macedonia':'MK','Macedonia, the Former Yugoslav Republic of':'MK','Swaziland':'SZ','Cape Verde':'CV',
 'Palestinian Territories':'PS','Palestine':'PS','Palestinian Authority':'PS','State of Palestine':'PS','Vatican City':'VA',
 'Holy See':'VA','United States of America':'US','Taiwan':'TW','Taiwan (China)':'TW','China, People\'s Republic of':'CN',
 'Kyrgyz Republic':'KG','Slovak Republic':'SK','Bosnia & Herzegovina':'BA','Bosnia-Herzegovina':'BA','East Timor':'TL',
 'St. Kitts and Nevis':'KN','St. Lucia':'LC','St. Vincent and the Grenadines':'VC','Sao Tome & Principe':'ST',
 'Republic of the Congo':'CG','Brunei':'BN','Brunei Darussalam':'BN','U.S. Virgin Islands':'VI','Virgin Islands (U.S.)':'VI',
 'British Virgin Islands':'VG','Virgin Islands (British)':'VG','Curacao':'CW','Bonaire, Sint Eustatius, and Saba':'BQ',
 'Republic of Moldova':'MD','Democratic People’s Republic of Korea':'KP','Lao People’s Democratic Republic':'LA',
 'United Republic of Tanzania':'TZ','United Kingdom of Great Britain and Northern Ireland':'GB',
 'Svalbard and Jan Mayen Islands':'SJ','Netherlands Antilles':'extra-netherlands-antilles',
 'England':'GB','Scotland':'GB','Wales':'GB','Northern Ireland':'GB',
 'West Bank':'extra-west-bank','Kosovo':'XK','European Union':'extra-european-union',
}
for n,k in aliases.items():names[norm(n)]=k
def entity(label):
 label=str(label or 'Unspecified').strip();key=names.get(norm(label))
 if not key:key='extra-'+re.sub(r'[^a-z0-9]+','-',norm(label))[:70]
 if key not in entities:entities[key]={'id':key,'name':label.title() if label.isupper() else label,'kind':'source-category','iso3':None,'aliases':[]}
 if label not in entities[key]['aliases']:entities[key]['aliases'].append(label)
 return key
s=BeautifulSoup((RAW/'un-members.html').read_text(encoding='utf-8'),'html.parser')
un=[h.get_text(' ',strip=True) for h in s.select('h2') if h.get_text(' ',strip=True)!='Search the United Nations']
unmatched=[]
for name in un:
 k=names.get(norm(name))
 if not k:unmatched.append(name)
 else:entities[k]['kind']='un-member'
for k in ['PS','VA']:entities[k]['kind']='un-observer'
entities['XK']={'id':'XK','name':'Kosovo','kind':'additional-jurisdiction','iso3':None,'aliases':['Kosovo']}
for k in ['TW','CK','NU']:entities[k]['kind']='additional-jurisdiction'
print('UN',len(un),'unmatched',unmatched,flush=True)
records=defaultdict(list);metrics=defaultdict(lambda:defaultdict(lambda:defaultdict(lambda:{'amount':0,'count':0})))
def add(r):records[r['country']].append(r)
def metric(k,m,y,amt):
 obj=metrics[k][m][str(y)];obj['count']+=1
 if amt is not None:obj['amount']+=amt

# One fixed ED release. Contract amounts are assigned to commencement, gifts to receipt.
sid=source('section117-2025-02-28.xlsx','government-data','U.S. Department of Education · Section 117 · 28 February 2025','Source snapshot through 2025-02-28; event years 2006–2025 selected; 2025 incomplete')
w=openpyxl.load_workbook(RAW/'section117-2025-02-28.xlsx',read_only=True,data_only=True);it=w.active.values;next(it);headers=next(it)
edstats=Counter();eddup=Counter();edrows=[];excluded=Counter()
for rn,row in enumerate(it,3):
 typ=row[3];contract='Contract' in str(typ);gift='Gift' in str(typ)
 receipt=date(row[7]);start=date(row[8]);end=date(row[9]);basis='contract-start' if contract and start else 'receipt'
 dt=start if contract and start else receipt;year=int(dt[:4]) if dt else None
 if year and not 2006<=year<=2025:excluded[str(year)]+=1;continue
 k=entity(row[5]);amount=number(row[6]);desc=str(row[12] or '')
 government=row[4] if row[4] in ['Yes','No'] else None
 payer=str(row[10] or row[11] or '') or None
 m='university-contract' if contract else 'university-gift' if gift else 'university-other'
 r={'id':f'ed-r{rn}','country':k,'origin':row[5],'category':'university','measure':m,'amount':amount,'year':year,'date':dt,'dateBasis':basis if dt else 'undated','periodStart':start,'periodEnd':end,'recipient':row[1],'principal':payer,'governmentSource':government,'type':typ,'description':desc[:1400] or None,'source':sid,'locator':f'Sheet1!A{rn}:R{rn}','legacy':row[17]=='Yes','cashVerified':False,'useLocation':'not-separable','review':'source-row'}
 # A textual campus indicator is evidence about a description, not a geographical split of the amount.
 if re.search(r'campus|WCM/Q|Weill Cornell.*Qatar|Education City',desc,re.I):r['useLocation']='campus-mentioned'
 dupkey=tuple(str(x) for x in row);eddup[dupkey]+=1
 if eddup[dupkey]>1:r['possibleDuplicate']=True
 add(r);edrows.append(r);edstats[m]+=1
 if year:
  metric(k,m,year,amount)
  if not r.get('possibleDuplicate'):metric(k,m+'-distinct',year,amount)
print('ED',len(edrows),dict(edstats),'undated',sum(r['year'] is None for r in edrows),flush=True)

# FARA: individual six-month receipts; amounts never consolidated by country.
fara=json.loads((OUT/'fara.json').read_text(encoding='utf-8'));farastats=Counter()
for x in fara:
 k=entity(x['origin']);file=x['source']+'.pdf';sid=source(file,'government-report',f'DOJ FARA · {x["reportYear"]} · H{x["half"]}',f'Publication half-year {x["reportYear"]}-H{x["half"]}; receipt periods may differ')
 services=x['services'];media=bool(re.search(r'media|public relations|publicity|information|printed|broadcast',services,re.I));lobby=bool(re.search(r'lobby|policy|political',services,re.I))
 cat='lobby' if lobby else 'media' if media else 'representation'
 end=x['periodEnd'];year=int(end[:4]) if end else None
 r={'id':x['id'],'country':k,'origin':x['origin'],'category':cat,'measure':'fara-receipts','amount':x['amount'],'year':year,'periodEnd':end,'periodMonths':x['periodMonths'],'reportYear':x['reportYear'],'reportHalf':x['half'],'recipient':x['recipient'],'principal':x['principal'],'governmentSource':None,'services':services,'description':x['excerpt'],'source':sid,'locator':f'PDF p. {x["page"]}; amount p. {x["amountPage"]}','page':x['page'],'amountPage':x['amountPage'],'registration':x['registration'],'duplicateOf':x.get('duplicateOf'),'review':x['review'],'periodText':x['finance'],'cashVerified':False}
 if year==2006 and end<'2006-06-30':r['boundaryOverlap']=True
 if lobby and media:r['mixedServices']=True
 add(r);farastats[cat]+=1
 if year and r['amount'] is not None and not r['duplicateOf']:metric(k,'fara-records',year,None)
print('FARA',len(fara),dict(farastats),flush=True)

# BEA historical-cost stocks and net financial transactions, parent and UBO bases.
files=[p for p in RAW.glob('bea-*') if ('detailed-country' in p.name.lower() and 'position' not in p.name.lower()) or 'Detailed%20Country' in p.name or 'ubo-detailed-country-position' in p.name]
beastats=Counter()
for p in files:
 w=openpyxl.load_workbook(p,read_only=True,data_only=True);ubo='ubo' in p.name.lower();sid=source(p.name,'government-statistic','BEA · '+p.name[4:],'2006–2025; historical-cost basis; values in source USD million')
 for sh in w:
  if sh.title.lower()=='notes':continue
  legacy='Detailed%20Country' in p.name
  measure='fdi-ubo-stock' if ubo else 'fdi-parent-stock' if sh.title=='Position' or legacy else 'fdi-net-flows' if 'Financial' in sh.title else None
  if not measure:continue
  vals=list(sh.values);yrrow=next((i for i,row in enumerate(vals[:12]) if sum(isinstance(v,(int,float)) and 1980<=v<=2025 for v in row)>2),None)
  if yrrow is None:continue
  cols={j:int(v) for j,v in enumerate(vals[yrrow]) if isinstance(v,(int,float)) and 2006<=v<=2025}
  colmeasure={j:measure for j in cols}
  if legacy:
   current=None
   for j,v in enumerate(vals[yrrow-1]):
    if isinstance(v,str) and 'Position' in v:current='fdi-parent-stock'
    elif isinstance(v,str) and 'Financial' in v:current='fdi-net-flows'
    elif isinstance(v,str) and 'Income' in v:current=None
    if j in colmeasure:colmeasure[j]=current
  region='Unspecified region'
  for rn,row in enumerate(vals[yrrow+1:],yrrow+2):
   label=row[0]
   if not label or not isinstance(label,str) or len(label)>110:continue
   if not any(row[j] is not None and row[j]!='' for j in cols):continue
   cleanlabel=re.sub(r'\s+\d+$','',label.strip())
   if cleanlabel in ['Europe','Africa','Middle East','Asia and Pacific','Latin America and Other Western Hemisphere','South and Central America','Other Western Hemisphere']:region=cleanlabel
   if cleanlabel=='Other':cleanlabel='Other · '+region
   k=entity(cleanlabel)
   for j,y in cols.items():
    measure=colmeasure[j]
    if measure is None:continue
    raw=row[j];v=number(raw);amt=v*1000000 if v is not None else None
    r={'id':f'{sid}-{sh.title}-{rn}-{j}-{y}','country':k,'origin':label.strip(),'category':'investment','measure':measure,'amount':amt,'sourceValue':raw,'year':y,'periodEnd':f'{y}-12-31','recipient':None,'principal':None,'governmentSource':None,'source':sid,'locator':f'{sh.title}!{openpyxl.utils.get_column_letter(j+1)}{rn}','review':'source-cell','cashVerified':False,'status':'reported' if v is not None else 'suppressed-or-range'}
    add(r);beastats[measure]+=1
    if amt is not None:metric(k,measure,y,amt)
print('BEA',dict(beastats),flush=True)

# Secondary tracker is kept as discovery evidence, never promoted to primary payment evidence.
for p in RAW.glob('think-*.html'):
 if p.name=='think-donors.html' or p.name not in mf:continue
 soup=BeautifulSoup(p.read_text(encoding='utf-8'),'html.parser');hh=soup.select('h1');h=hh[-1] if hh else None;table=soup.select_one('table.single-donor')
 if not h or not table:continue
 k=entity(h.get_text(' ',strip=True));sid=source(p.name,'secondary-tracker','Think Tank Funding Tracker · '+entities[k]['name'],'2019–2024 aggregate; donor-year primary documents not all checked')
 for n,tr in enumerate(table.select('tbody tr'),1):
  td=tr.select('td')
  if len(td)<3:continue
  val=number(td[2].get_text(' ',strip=True));unknown=bool(td[2].select('.not-disclosed'))
  if unknown:val=None
  add({'id':f'{sid}-{n}','country':k,'origin':h.get_text(' ',strip=True),'category':'thinktank','measure':'tracker-minimum','amount':val,'year':None,'periodStart':'2019-01-01','periodEnd':'2024-12-31','recipient':td[0].get_text(' ',strip=True),'principal':td[1].get_text(' ',strip=True),'source':sid,'locator':f'Donor table row {n}','review':'secondary-unverified','cashVerified':False,'governmentSource':'tracker classification','uncertainty':'Aggregated disclosure lower bounds, not a verified payment; year allocation and original documents pending'})

# Domestic US organizations, kept under US and never attributed to a foreign country.
irs=[]
for p in RAW.glob('irs-render-*.html'):
 if p.name not in mf:continue
 s=BeautifulSoup(p.read_bytes(),'html.parser')
 def value(field):
  el=s.find(id=re.compile('/'+field+r'\[1\]$'));return el.get_text(' ',strip=True) if el else None
 end=value('TaxPeriodEndDt');start=value('TaxPeriodBeginDt');exp=number(value('CYTotalExpensesAmt'));rev=number(value('CYTotalRevenueAmt'))
 if not end:
  txt=s.get_text(' ',strip=True);m=re.search(r'tax year beginning\s+(\d\d-\d\d-\d{4})\s*,?\s*and ending\s+(\d\d-\d\d-\d{4})',txt)
  if m:start,end=[datetime.datetime.strptime(v,'%m-%d-%Y').date().isoformat() for v in m.groups()]
 if not end or exp is None:continue
 if re.fullmatch(r'\d\d-\d\d-\d{4}',end):end=datetime.datetime.strptime(end,'%m-%d-%Y').date().isoformat()
 if start and re.fullmatch(r'\d\d-\d\d-\d{4}',start):start=datetime.datetime.strptime(start,'%m-%d-%Y').date().isoformat()
 sid=source(p.name,'primary-filing-mirror','AIPAC · IRS Form 990 · FY ending '+end,'Original filing rendered by ProPublica; Part I lines 12 and 18; fiscal year')
 r={'id':sid,'country':'US','origin':'United States','category':'domestic','measure':'organization-expenses','amount':exp,'revenue':rev,'year':int(end[:4]),'periodStart':start,'periodEnd':end,'recipient':'AIPAC · organization expenses (not a named payee)','principal':'American Israel Public Affairs Committee · EIN 53-0217164','source':sid,'locator':'Form 990, Part I, line 18, Current Year','review':'source-field','cashVerified':False}
 add(r);irs.append(r)
print('IRS',[(r['year'],r['amount']) for r in irs],flush=True)
fec=[]
for p in RAW.glob('fec-committee-summary-*.zip'):
 cycle=int(p.stem[-4:]);sid=source(p.name,'government-data',f'FEC · PAC and party summary · {cycle} cycle',f'{cycle-1}–{cycle}; coverage end dates vary; includes 2005 in the 2006 cycle')
 with zipfile.ZipFile(p) as z:
  file=z.namelist()[0];lines=z.read(file).decode('utf-8',errors='replace').splitlines()
 for rn,line in enumerate(lines,1):
  a=line.split('|');name=a[1]
  if not re.search(r'AMERICAN ISRAEL PUBLIC AFFAIRS|UNITED DEMOCRACY PROJECT|US.?INDIA POLITICAL|U\.?S\.?-?CUBA DEMOCRACY|ARMENIAN NATIONAL|JSTREETPAC',name,re.I):continue
  end=datetime.datetime.strptime(a[26],'%m/%d/%Y').date().isoformat() if a[26] else None
  r={'id':f'fec-{a[0]}-{cycle}','country':'US','origin':'United States','category':'domestic','measure':'committee-disbursements','amount':number(a[12]),'receipts':number(a[5]),'independentExpenditures':number(a[23]),'affiliateTransfers':number(a[13]),'year':cycle,'periodStart':f'{cycle-1}-01-01','periodEnd':end,'recipient':name+' · committee disbursements (not a named payee)','principal':name,'committeeId':a[0],'source':sid,'locator':f'{file}, line {rn}; CMTE_ID {a[0]}; TTL_DISB','review':'source-row','cashVerified':False,'boundaryOverlap':cycle==2006}
  add(r);fec.append(r)
print('FEC',len(fec),[(r['committeeId'],r['principal'],r['amount']) for r in fec if r['year']==2024],flush=True)

# Manually reviewed original think-tank donor rolls, if present.
primary=OUT/'think-primary.json'
if primary.exists():
 for r in json.loads(primary.read_text(encoding='utf-8')):
  r['origin']=r.pop('originName');r['country']=entity(r['origin']);r['source']=source(r.pop('sourceFile'),'primary-recipient-report',r.pop('sourceTitle'),r.get('periodStart','')+' – '+r.get('periodEnd',''));add(r)

catalog=[]
for k,e in entities.items():
 rr=records[k];counts=Counter(r['category'] for r in rr);primaryCounts=Counter(r['category'] for r in rr if r['review']!='secondary-unverified');availableYears=sorted({r['year'] for r in rr if r.get('year') and r['review']!='secondary-unverified'})
 e.update({'recordCount':len(rr),'counts':dict(counts),'primaryCounts':dict(primaryCounts),'years':availableYears,'metrics':metrics[k]})
 catalog.append(e);write(DEST/(k+'.json'),rr)
catalog.sort(key=lambda e:e['name']);write(DEST/'countries.json',catalog);write(DEST/'sources.json',list(sources.values()))
stats={'asOf':'2026-10-26','targetPeriod':[2006,2025],'entities':len(entities),'unMembers':sum(e['kind']=='un-member' for e in catalog),'unObservers':2,'records':sum(len(x) for x in records.values()),'edRows':len(edrows),'edUndated':sum(r['year'] is None for r in edrows),'edPotentialDuplicateRows':sum(r.get('possibleDuplicate',False) for r in edrows),'faraRows':len(fara),'faraAmounts':sum(r['amount'] is not None for r in fara),'faraSources':40,'sourceCount':len(sources),'edExcludedYears':dict(excluded),'faraConsolidatedSpending':None,'coverageComplete':False,'unknownUNMappings':unmatched}
write(DEST/'meta.json',stats);write(OUT/'dataset-stats.json',stats)
# Export self-contained line-delimited evidence, catalogue and source register in one compressed download.
with zipfile.ZipFile(DEST/'evidence.zip','w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
 for name in ['countries.json','sources.json','meta.json']:z.write(DEST/name,name)
 with z.open('evidence.jsonl','w') as f:
  for rr in records.values():
   for r in rr:f.write((json.dumps(r,ensure_ascii=False)+'\n').encode())
print('DONE',stats,flush=True)
